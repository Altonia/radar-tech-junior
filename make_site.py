"""Génère le site GitHub Pages (docs/data.json) et l'Excel (docs/m13.xlsx)."""
import json, os, re, datetime, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from common import ROOT

q = lambda u: urllib.parse.quote(u, safe=":/?&=#%+,@;~!$'()*[]") if u else u
P = lambda f: os.path.join(ROOT, f)
REG = {"Île-de-France": "idf", "Pays de la Loire": "pdl", "Bretagne": "bzh"}
today = datetime.date.today()

rows = json.load(open(P("data/annuaire.json")))
libres = json.load(open(P("data/offres_libres.json"))) if os.path.exists(P("data/offres_libres.json")) else []

allo = [dict(x, _ent=r["nom"], _site=r["site"] if r.get("site_niveau") == "verifie" else None) for r in rows for x in r["offres"]]
allo += [dict(x, _ent=x.get("entreprise"), _site=None) for x in libres]
seen, offres = set(), []
for x in sorted(allo, key=lambda x: x.get("date") or "", reverse=True):
    if str(x["id"]) in seen: continue
    seen.add(str(x["id"]))
    offres.append({"id": str(x["id"]), "t": x["titre"], "f": x["famille"], "al": bool(x.get("alternance")), "k": x.get("contrat"),
                   "l": x.get("lieu"), "d": (x.get("date") or "")[:10], "u": q(x.get("url")), "e": x["_ent"],
                   "w": q(x["_site"]), "r": x["region"], "src": x["source"], "cab": bool(x.get("via_cabinet"))})

ents = []
for i, r in enumerate(rows):
    d = {"id": f'{r["siren"] or "x" + str(i)}-{REG[r["region"]]}', "n": r["nom"], "s": r["segment"], "r": r["region"],
         "v": r["ville"], "c": r["cp"], "e": r["effectif"], "tr": r["tranche"], "a": r["naf"], "sc": r["score_ent"]}
    if r["siren"]: d["si"] = r["siren"]
    tags = [t for t in r["tags"] if t != r["segment"]]
    if tags: d["t"] = tags
    if r["site"]: d["w"] = q(r["site"]); d["wp"] = r.get("site_niveau") == "probable"
    if r.get("lba") and r["lba"].get("url"): d["lba"] = q(r["lba"]["url"])
    if r["lbb"]: d["pot"] = 1
    # vrais dirigeants opérationnels seulement (pas les commissaires aux comptes, administrateurs, liquidateurs…)
    def rang(role):
        rl = role.lower()
        if re.search(r"commissaire|liquidat|membre du conseil|contr[ôo]leur|^administrateur$|^autre$|^membre$", rl) or not rl: return None
        for i, k in enumerate(["directeur général", "président", "gérant", "directeur général délégué", "directoire", "dirigeant", "pouvoir d’engager"]):
            if k in rl: return i
        return None
    dg = sorted(((rang(x["role"]), x) for x in (r.get("dirigeants") or [])), key=lambda t: 99 if t[0] is None else t[0])
    dg = [[f'{x["prenom"]} {x["nom"]}'.strip(), x["role"]] for rk, x in dg if rk is not None][:2]
    if dg: d["dg"] = dg
    if r["offres"]: d["o"] = [str(x["id"]) for x in r["offres"]]
    ents.append(d)
ents.sort(key=lambda d: -d["sc"])
meta = {"maj": today.strftime("%d/%m/%Y"), "ref": today.isoformat(), "familles": sorted({x["f"] for x in offres})}
os.makedirs(P("docs"), exist_ok=True)
json.dump({"meta": meta, "entreprises": ents, "offres": offres}, open(P("docs/data.json"), "w"), ensure_ascii=False, separators=(",", ":"))

# --- Excel ---
wb = Workbook()
HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="2B45C9"); LINK = Font(color="2B45C9", underline="single")

def sheet(ws, cols, data, links=()):
    ws.append([c for c, _ in cols])
    for c in ws[1]: c.font, c.fill, c.alignment = HEAD, FILL, Alignment(vertical="center")
    for row in data: ws.append(row)
    for j, (_, w) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(j)].width = w
    for j in links:
        for cell in ws.iter_rows(min_row=2, min_col=j, max_col=j):
            c = cell[0]
            if c.value: c.hyperlink, c.font = c.value, LINK
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions

nb_off = {d["id"]: len(d.get("o", [])) for d in ents}
ws = wb.active; ws.title = "Entreprises"
sheet(ws, [("Entreprise", 34), ("Type d'entreprise", 28), ("Région", 16), ("Ville", 20), ("Code postal", 11), ("Salariés", 11),
           ("Activité", 34), ("Prend des alternants", 12), ("Potentiel d'embauche", 12), ("Offres en cours", 10), ("Pertinence", 10),
           ("Site", 36), ("Site vérifié", 10), ("Candidater en alternance", 36), ("SIREN", 12), ("Fiche officielle", 36), ("Candidaté (à cocher)", 12)],
      [[d["n"], d["s"], d["r"], d["v"], d["c"], d["e"], d["a"], "oui" if d.get("lba") else "", "oui" if d.get("pot") else "",
        nb_off[d["id"]] or "", d["sc"], d.get("w"), ("oui" if not d.get("wp") else "probable") if d.get("w") else "", d.get("lba"),
        d.get("si"), f'https://annuaire-entreprises.data.gouv.fr/entreprise/{d["si"]}' if d.get("si") else None, ""] for d in ents],
      links=(12, 14, 16))
ws = wb.create_sheet("Offres")
sheet(ws, [("Intitulé", 50), ("Entreprise", 30), ("Type de poste", 22), ("Contrat", 14), ("Lieu", 24), ("Région", 16),
           ("Publiée le", 12), ("Source", 16), ("Lien de l'offre", 40), ("Site de l'entreprise", 34), ("Candidaté (à cocher)", 12)],
      [[x["t"], x["e"] or "Non précisée", x["f"], "Alternance" if x["al"] else ("CDI" if kind == "cdi" else x["k"] or ""), x["l"], x["r"],
        x["d"], x["src"], x["u"], x["w"], ""] for x in offres for kind in [("cdi" if (x["k"] or "").lower() in ("cdi", "permanent") else "")]],
      links=(9, 10))
ws = wb.create_sheet("Mode d'emploi")
for line in [
    "M13 — mis à jour le " + meta["maj"],
    "",
    "Onglet Entreprises : pour les candidatures spontanées. Utilisez les flèches de filtre en haut de chaque colonne.",
    "« Prend des alternants » : l'entreprise recrute régulièrement des alternants (La Bonne Alternance).",
    "« Potentiel d'embauche » : forte probabilité de recrutement dans ces métiers (La Bonne Boîte, France Travail).",
    "« Site vérifié » = oui : le SIREN de l'entreprise figure sur le site. « probable » : le domaine correspond au nom, sans preuve formelle.",
    "« Pertinence » : ordre de tri (type d'entreprise, taille, signaux de recrutement), pas un avis sur l'entreprise.",
    "",
    "Onglet Offres : offres ouvertes aux débutants (2 ans d'expérience maximum) et alternances, publiées depuis 31 jours au plus.",
    "Déjà retiré : postes seniors, écoles qui recrutent des étudiants, niveaux BTS ou Bac+3, missions freelance, doublons, liens morts.",
    "Conseil : quand c'est possible, postulez sur le site carrières de l'entreprise plutôt que via l'agrégateur.",
    "",
    "Sources : base SIRENE, France Travail, La Bonne Alternance, Adzuna, partenaires Odoo, Wikidata.",
]:
    ws.append([line])
ws.column_dimensions["A"].width = 130; ws["A1"].font = Font(bold=True, size=14)
wb.save(P("docs/m13.xlsx"))
print(len(ents), "entreprises,", len(offres), "offres →", round(os.path.getsize(P("docs/data.json")) / 1e6, 2), "Mo JSON,",
      round(os.path.getsize(P("docs/m13.xlsx")) / 1e6, 2), "Mo Excel")
